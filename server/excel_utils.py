from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="0B8CFF")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2EC"),
    right=Side(style="thin", color="D9E2EC"),
    top=Side(style="thin", color="D9E2EC"),
    bottom=Side(style="thin", color="D9E2EC"),
)


def ler_csv(caminho: Path, campos: list[str]) -> list[dict[str, str]]:
    if not caminho.exists():
        return []

    with caminho.open("r", encoding="utf-8-sig", newline="") as arquivo:
        leitor = csv.DictReader(arquivo, delimiter=";")
        linhas = []
        for linha in leitor:
            linhas.append({campo: (linha.get(campo) or "") for campo in campos})
        return linhas


def ajustar_colunas(ws, campos: list[str]) -> None:
    largura_padrao = {
        "url": 60,
        "descricao": 70,
        "localizacao": 45,
        "anunciante": 32,
        "data_anuncio": 42,
        "preco": 16,
        "metragem": 14,
        "quartos": 12,
        "data_coleta": 14,
        "hora_coleta": 14,
    }

    for idx, campo in enumerate(campos, start=1):
        letra = get_column_letter(idx)
        if campo in largura_padrao:
            ws.column_dimensions[letra].width = largura_padrao[campo]
        else:
            maior = len(campo)
            for cell in ws[letra][1:200]:
                maior = max(maior, len(str(cell.value or "")))
            ws.column_dimensions[letra].width = min(max(maior + 2, 12), 35)


def criar_aba(wb: Workbook, nome: str, campos: list[str], linhas: list[dict[str, str]], tabela_nome: str) -> None:
    ws = wb.create_sheet(nome)
    ws.append(campos)

    for linha in linhas:
        ws.append([linha.get(campo, "") for campo in campos])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER

    ajustar_colunas(ws, campos)

    if ws.max_row >= 2 and ws.max_column >= 1:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tabela = Table(displayName=tabela_nome, ref=ref)
        estilo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tabela.tableStyleInfo = estilo
        ws.add_table(tabela)


def gerar_excel_final(
    arquivo_lista_csv: Path,
    arquivo_detalhe_csv: Path,
    arquivo_excel_final: Path,
    campos_lista: list[str],
    campos_detalhe: list[str],
) -> dict[str, int | str]:
    arquivo_excel_final.parent.mkdir(parents=True, exist_ok=True)

    linhas_lista = ler_csv(arquivo_lista_csv, campos_lista)
    linhas_detalhe = ler_csv(arquivo_detalhe_csv, campos_detalhe)

    wb = Workbook()
    aba_padrao = wb.active
    wb.remove(aba_padrao)

    criar_aba(wb, "Anuncios_Detalhes", campos_detalhe, linhas_detalhe, "TabelaDetalhes")
    criar_aba(wb, "URLs_Listagem", campos_lista, linhas_lista, "TabelaURLs")

    # Aba de resumo simples.
    ws = wb.create_sheet("Resumo", 0)
    ws.append(["Métrica", "Valor"])
    ws.append(["Cidade", "Presidente Prudente - SP"])
    ws.append(["Tipo", "Compra / Venda"])
    ws.append(["Total de anúncios detalhados", len(linhas_detalhe)])
    ws.append(["Total de URLs coletadas na listagem", len(linhas_lista)])
    ws.append(["Arquivo de detalhes CSV", str(arquivo_detalhe_csv)])
    ws.append(["Arquivo de URLs CSV", str(arquivo_lista_csv)])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    ws.freeze_panes = "A2"

    wb.save(arquivo_excel_final)

    return {
        "total_lista": len(linhas_lista),
        "total_detalhes": len(linhas_detalhe),
        "arquivo_excel_final": str(arquivo_excel_final),
    }
