from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path
from urllib.parse import unquote

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TIPOS = {
    "apartamento": ("Apartamento", ("apartamento",)),
    "cobertura": ("Cobertura", ("cobertura",)),
    "loft": ("Loft", ("loft",)),
    "studio": ("Studio", ("studio",)),
    "estudio": ("Studio", ("studio",)),
    "kitnet": ("Kitnet/Quitinete", ("kitnet", "quitinete", "conjugado")),
    "quitinete": ("Kitnet/Quitinete", ("kitnet", "quitinete", "conjugado")),
    "kitnet conjugado": ("Kitnet/Quitinete", ("kitnet", "quitinete", "conjugado")),
    "flat": ("Flat", ("flat",)),
    "casa": ("Casa", ("casa",)),
    "casa de condominio": ("Casa de Condomínio", ("casa-de-condominio", "casa-condominio")),
    "casa em condominio": ("Casa de Condomínio", ("casa-de-condominio", "casa-condominio")),
    "casa de vila": ("Casa de Vila", ("casa-de-vila",)),
    "sobrado": ("Sobrado", ("sobrado",)),
    "terreno": ("Terreno/Lote", ("terreno", "lote")),
    "lote": ("Terreno/Lote", ("terreno", "lote")),
    "lote terreno": ("Terreno/Lote", ("terreno", "lote")),
    "terreno lote": ("Terreno/Lote", ("terreno", "lote")),
    "lote em condominio": ("Terreno/Lote em Condomínio", ("lote-em-condominio", "terreno-em-condominio")),
    "terreno em condominio": ("Terreno/Lote em Condomínio", ("lote-em-condominio", "terreno-em-condominio")),
}


def normalizar(valor: object) -> str:
    if pd.isna(valor):
        return ""
    texto = unicodedata.normalize("NFD", str(valor)).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[_/|-]+", " ", texto.lower())
    return re.sub(r"\s+", " ", texto).strip()


def url_normalizada(valor: object) -> str:
    if pd.isna(valor):
        return ""
    return unquote(str(valor)).lower()


def inferir_tipo_url(url: object) -> str:
    caminho = url_normalizada(url)
    # Específicos primeiro para não classificar casa de condomínio apenas como casa.
    regras = (
        ("Casa de Condomínio", ("casa-de-condominio", "casa-condominio")),
        ("Casa de Vila", ("casa-de-vila",)),
        ("Terreno/Lote em Condomínio", ("lote-em-condominio", "terreno-em-condominio")),
        ("Cobertura", ("venda-cobertura", "aluguel-cobertura")),
        ("Apartamento", ("venda-apartamento", "aluguel-apartamento")),
        ("Loft", ("venda-loft", "aluguel-loft")),
        ("Studio", ("venda-studio", "aluguel-studio")),
        ("Kitnet/Quitinete", ("venda-kitnet", "aluguel-kitnet", "venda-quitinete", "venda-conjugado")),
        ("Flat", ("venda-flat", "aluguel-flat")),
        ("Sobrado", ("venda-sobrado", "aluguel-sobrado")),
        ("Casa", ("venda-casa", "aluguel-casa")),
        ("Terreno/Lote", ("venda-terreno", "aluguel-terreno", "venda-lote", "aluguel-lote")),
    )
    for tipo, termos in regras:
        if any(termo in caminho for termo in termos):
            return tipo
    return ""


def canonizar_tipo(valor: object) -> tuple[str, tuple[str, ...]]:
    chave = normalizar(valor)
    if chave in TIPOS:
        return TIPOS[chave]
    for candidato, resultado in TIPOS.items():
        if candidato and candidato in chave:
            return resultado
    return "", ()


def extrair_id(url: object) -> str:
    if pd.isna(url):
        return ""
    match = re.search(r"(?:^|[-_/])id[-_]?([0-9]+)(?:[/?#]|$)", str(url), re.I)
    return match.group(1) if match else ""


def escolher_aba(arquivo: Path) -> str:
    excel = pd.ExcelFile(arquivo)
    preferidas = ("Todos os imóveis", "Todos os imoveis", "Imóveis", "Imoveis", "Dados")
    for nome in preferidas:
        if nome in excel.sheet_names:
            return nome
    for nome in excel.sheet_names:
        colunas = [normalizar(c) for c in pd.read_excel(arquivo, sheet_name=nome, nrows=0).columns]
        if "url" in colunas or "id anuncio" in colunas:
            return nome
    raise ValueError(f"Nenhuma aba de imóveis reconhecida em {arquivo.name}. Abas: {excel.sheet_names}")


def carregar(arquivo: Path) -> pd.DataFrame:
    aba = escolher_aba(arquivo)
    df = pd.read_excel(arquivo, sheet_name=aba, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    aliases = {normalizar(c): c for c in df.columns}
    renomear = {}
    for destino, opcoes in {
        "id_anuncio": ("id anuncio", "id", "codigo anuncio", "codigo do anuncio"),
        "url": ("url", "link", "url anuncio", "link anuncio"),
        "tipo_imovel": ("tipo imovel", "tipo do imovel", "tipo"),
    }.items():
        if destino not in df.columns:
            origem = next((aliases[x] for x in opcoes if x in aliases), None)
            if origem:
                renomear[origem] = destino
    df = df.rename(columns=renomear)
    if "url" not in df.columns:
        raise ValueError(f"A aba {aba!r} de {arquivo.name} não possui coluna URL/link.")
    if "id_anuncio" not in df.columns:
        df["id_anuncio"] = ""
    if "tipo_imovel" not in df.columns:
        df["tipo_imovel"] = ""
    df["arquivo_origem"] = arquivo.name
    df["aba_origem"] = aba
    df["linha_origem"] = range(2, len(df) + 2)
    return df


def auditar_tipos(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    ids = df["id_anuncio"].fillna("").astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    extraidos = df["url"].map(extrair_id)
    df["id_anuncio"] = ids.mask(ids.eq(""), extraidos)
    df["tipo_imovel_original"] = df["tipo_imovel"]

    tratados, inferidos, status, motivos = [], [], [], []
    for _, row in df.iterrows():
        canonico, termos = canonizar_tipo(row.get("tipo_imovel"))
        inferido = inferir_tipo_url(row.get("url"))
        url = url_normalizada(row.get("url"))
        motivo = ""
        if not str(row.get("id_anuncio") or "").strip():
            motivo = "ID do anúncio ausente e não encontrado na URL"
        elif canonico and termos and not any(t in url for t in termos):
            motivo = f"Tipo informado '{canonico}' não aparece na URL"
        elif canonico and inferido and canonico != inferido:
            motivo = f"Tipo informado '{canonico}' diverge do tipo inferido pela URL '{inferido}'"
        elif not canonico and inferido:
            canonico = inferido
        elif not canonico:
            motivo = "Tipo ausente/desconhecido e não inferido pela URL"
        tratados.append(canonico or "")
        inferidos.append(inferido)
        status.append("SUSPEITO" if motivo else "CONFIRMADO")
        motivos.append(motivo)

    df["tipo_imovel_tratado"] = tratados
    df["tipo_inferido_url"] = inferidos
    df["status_tipo"] = status
    df["motivo_suspeita"] = motivos
    return df


def deduplicar(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    df = df.copy()
    campos_dados = [c for c in df.columns if c not in {"arquivo_origem", "aba_origem", "linha_origem"}]
    df["_completude"] = df[campos_dados].notna().sum(axis=1) + df[campos_dados].astype(str).ne("").sum(axis=1)
    df["_ordem"] = range(len(df))
    com_id = df["id_anuncio"].fillna("").astype(str).str.strip().ne("")
    ordenado = df.sort_values(["id_anuncio", "_completude", "_ordem"], ascending=[True, False, True])
    repetido = com_id.loc[ordenado.index] & ordenado.duplicated("id_anuncio", keep="first")
    removidos = ordenado[repetido].copy()
    unicos = pd.concat([ordenado[~repetido]], ignore_index=True).sort_values("_ordem")
    return unicos.drop(columns=["_completude", "_ordem"]), removidos.drop(columns=["_completude", "_ordem"])


def formatar_excel(writer: pd.ExcelWriter) -> None:
    azul, amarelo, vermelho = "1F4E78", "FFF2CC", "F4CCCC"
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=azul)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for col in range(1, ws.max_column + 1):
            valores = [len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 200) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(55, max(11, max(valores, default=10) + 2))
        cabecalhos = {str(c.value): c.column for c in ws[1]}
        if "status_tipo" in cabecalhos and ws.max_row >= 2:
            col = get_column_letter(cabecalhos["status_tipo"])
            faixa = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
            ws.conditional_formatting.add(faixa, FormulaRule(formula=[f'${col}2="SUSPEITO"'], fill=PatternFill("solid", fgColor=vermelho)))
        if ws.title == "Resumo":
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, 1).value == "Tipos suspeitos":
                    ws.cell(row, 2).fill = PatternFill("solid", fgColor=amarelo)
                    ws.cell(row, 2).font = Font(bold=True)


def processar(arquivos: list[Path], saida: Path) -> dict[str, int]:
    origem = pd.concat([carregar(arquivo) for arquivo in arquivos], ignore_index=True, sort=False)
    auditado = auditar_tipos(origem)
    unicos, duplicados = deduplicar(auditado)
    suspeitos = unicos[unicos["status_tipo"].eq("SUSPEITO")].copy()
    contagem = (suspeitos.groupby(["tipo_imovel_original", "motivo_suspeita"], dropna=False)
                .size().reset_index(name="quantidade").sort_values("quantidade", ascending=False))
    resumo = pd.DataFrame([
        ("Linhas recebidas", len(origem)),
        ("Imóveis únicos", len(unicos)),
        ("Duplicatas removidas por ID", len(duplicados)),
        ("Tipos confirmados", int(unicos["status_tipo"].eq("CONFIRMADO").sum())),
        ("Tipos suspeitos", len(suspeitos)),
        ("Registros sem ID", int(unicos["id_anuncio"].fillna("").astype(str).str.strip().eq("").sum())),
    ], columns=["indicador", "quantidade"])

    saida.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(saida, engine="openpyxl") as writer:
        unicos.to_excel(writer, sheet_name="Imóveis unificados", index=False)
        suspeitos.to_excel(writer, sheet_name="Suspeitos", index=False)
        contagem.to_excel(writer, sheet_name="Contagem suspeitos", index=False)
        duplicados.to_excel(writer, sheet_name="Duplicatas removidas", index=False)
        resumo.to_excel(writer, sheet_name="Resumo", index=False)
        formatar_excel(writer)
    return {linha.indicador: int(linha.quantidade) for linha in resumo.itertuples(index=False)}


def main() -> int:
    raiz = Path(__file__).resolve().parent
    padrao = raiz / "dados" / "zap_v4_presidente_prudente_todos_tipos.xlsx"
    parser = argparse.ArgumentParser(description="Trata um ou une dois Excels do ZAP v4, deduplica por ID e audita o tipo pela URL.")
    parser.add_argument("excels", type=Path, nargs="*", help="Zero, um ou dois arquivos Excel. Sem argumentos, usa o consolidado em dados/.")
    parser.add_argument("-o", "--saida", type=Path, default=raiz / "dados" / "zap_v4_tratado_unificado.xlsx")
    args = parser.parse_args()
    if len(args.excels) > 2:
        parser.error("Informe no máximo dois arquivos Excel.")
    arquivos = args.excels or [padrao]
    for arquivo in arquivos:
        if not arquivo.is_file():
            if not args.excels:
                parser.error(
                    f"Excel consolidado não encontrado em: {arquivo}\n"
                    "Informe manualmente: py tratamento_unificar.py worker1.xlsx worker2.xlsx"
                )
            parser.error(f"Arquivo não encontrado: {arquivo}")
    try:
        resultado = processar([arquivo.resolve() for arquivo in arquivos], args.saida.resolve())
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    print(f"Arquivo gerado: {args.saida.resolve()}")
    print("Arquivo(s) de entrada:")
    for arquivo in arquivos:
        print(f"- {arquivo.resolve()}")
    for nome, valor in resultado.items():
        print(f"- {nome}: {valor}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
